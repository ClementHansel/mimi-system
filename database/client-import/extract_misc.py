"""Payroll (Gaji), daily sales by payment method, weekly outlet bills (supplier deliveries)."""
import openpyxl, glob, re, csv, sys, os, datetime, warnings
warnings.filterwarnings("ignore")
DATA = os.environ["CLIENT_DATA_DIR"]  # folder holding SO/, Gaji/, Pembukuan/, "Tagihan mingguan"/ (the client's Drive folders)
OUT = sys.argv[1]

def num(v):
    if v is None or v == "": return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip().replace(",", ""))
    except ValueError: return 0.0
def clean(s): return re.sub(r"\s+", " ", str(s)).strip() if s is not None else ""

# ---------- payroll
w = csv.writer(open(f"{OUT}/payroll.csv", "w", newline="", encoding="utf-8"))
cols = ["no","name","tenure","position","bank_account","work_days","overtime_hours","finger","cuti","sakit","izin","alpa","off",
        "gaji_pokok","lembur","insentif_performa","tunj_kehadiran","tunj_masa_kerja","tunj_jabatan","total_pendapatan",
        "pot_izin","pot_alpa","minus_barang","pinjaman","pot_finger","total_potongan","gaji_bersih"]
w.writerow(["file","unit","period","finance","address"] + cols)
for f in sorted(glob.glob(f"{DATA}/Gaji/*.xlsm")):
    base = os.path.basename(f)
    m = re.search(r"Mimi-(M\d+)", base)
    unit = m.group(1) if m else ("LOG_BPP" if "Logistik Bpp" in base else "LOG_STAFF")
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb["TABEL GAJI"]
    rows = list(ws.iter_rows(min_row=1, max_col=27, values_only=True))
    period = clean(rows[0][2]); finance = clean(rows[1][2]); address = clean(rows[4][2])
    for r in rows[8:]:
        if not isinstance(r[0], (int, float)) or not clean(r[1]): continue
        w.writerow([base, unit, period, finance, address] + [clean(v) if i in (1,2,3,4) else (r[i] if r[i] is not None else "") for i, v in enumerate(r[:27])])
    wb.close()

# ---------- sales by payment method
w = csv.writer(open(f"{OUT}/sales_payments.csv", "w", newline="", encoding="utf-8"))
METHODS = ["cash","shopeepay","shopeefood","grabfood","gofood","gopay","transfer","qris","voucher"]
w.writerow(["outlet","date"] + METHODS + ["total"])
wb = openpyxl.load_workbook(f"{DATA}/Pembukuan/Sales All Resto dan Store April 2026.xlsx", read_only=True, data_only=True)
for ws in wb.worksheets:
    if not re.fullmatch(r"M\d+", ws.title): continue
    for r in ws.iter_rows(min_row=5, max_col=11, values_only=True):
        if isinstance(r[0], datetime.datetime):
            w.writerow([ws.title, r[0].date().isoformat()] + [num(v) for v in r[1:11]])
wb.close()

# ---------- weekly outlet bills: line-level deliveries with supplier + price
w = csv.writer(open(f"{OUT}/outlet_bills.csv", "w", newline="", encoding="utf-8"))
w.writerow(["file","sheet","date","category","supplier","product","unit","qty","price","amount","note"])
for f in sorted(glob.glob(f"{DATA}/Tagihan mingguan/**/*OUTLET*.xlsx", recursive=True) + glob.glob(f"{DATA}/Tagihan mingguan/**/*OTLET*.xlsx", recursive=True)):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for ws in wb.worksheets:
        hdr = None; last_date = None
        for r in ws.iter_rows(min_row=1, max_col=20, values_only=True):
            cells = [clean(v).upper() for v in r]
            if hdr is None:
                if "NAMA SUPPLIER" in cells and "QTY" in cells:
                    hdr = {c: i for i, c in enumerate(cells) if c}
                continue
            get = lambda k: r[hdr[k]] if k in hdr else None
            prod = clean(get("NAMA PRODUCT") or get("NAMA PRODUK") or get("NAMA BARANG"))
            if not prod: continue
            dt = get("TANGGAL")
            if isinstance(dt, datetime.datetime): last_date = dt.date()
            w.writerow([os.path.basename(f), ws.title, last_date.isoformat() if last_date else "", clean(get("KATEGORI BARANG")),
                        clean(get("NAMA SUPPLIER")), prod, clean(get("SATUAN")).upper(), num(get("QTY")), num(get("HARGA")), num(get("JUMLAH")), clean(get("KETERANGAN"))])
    wb.close()
print("ok", file=sys.stderr)
